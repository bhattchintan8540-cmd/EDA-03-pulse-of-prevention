"""Build assignment-format PPTX and HTML decks from findings.json payloads."""
from __future__ import annotations

from pathlib import Path
import html

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .config import PRESENTATIONS, ROOT

NAVY = RGBColor(15, 23, 42)
TEAL = RGBColor(14, 116, 144)
WHITE = RGBColor(255, 255, 255)
SLATE = RGBColor(51, 65, 85)
LIGHT = RGBColor(241, 245, 249)


def _set_run(run, text, size=18, bold=False, color=NAVY):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _add_bar(slide, prs):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.18))
    shape.fill.solid()
    shape.fill.fore_color.rgb = TEAL
    shape.line.fill.background()
    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, prs.slide_height - Inches(0.42), prs.slide_width, Inches(0.42)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = NAVY
    footer.line.fill.background()


def _blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bar(slide, prs)
    return slide


def _title_box(slide, text, top=0.35, size=28):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    _set_run(p.add_run(), text, size=size, bold=True, color=NAVY)
    return box


def _bullets(slide, items, top=1.15, height=5.4):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12.0), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        _set_run(p.add_run(), f"•  {item}", size=18, color=SLATE)


def _as_lines(value, limit=8):
    if isinstance(value, list):
        return [str(v) for v in value[:limit]]
    if isinstance(value, dict):
        lines = []
        for k, v in list(value.items())[:limit]:
            lines.append(f"{k}: {v}")
        return lines
    return [str(value)]


def build_pptx(findings: dict) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Title
    slide = _blank(prs)
    fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fill.fill.solid()
    fill.fill.fore_color.rgb = NAVY
    fill.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.28), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2))
    _set_run(box.text_frame.paragraphs[0].add_run(), findings["code"], size=18, bold=True, color=TEAL)
    box2 = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), Inches(11.5), Inches(2))
    _set_run(box2.text_frame.paragraphs[0].add_run(), findings["title"], size=32, bold=True, color=WHITE)
    box3 = slide.shapes.add_textbox(Inches(0.9), Inches(5.2), Inches(11.5), Inches(1))
    _set_run(box3.text_frame.paragraphs[0].add_run(), findings.get("dataset", ""), size=16, color=LIGHT)

    # 1 Overview
    s = _blank(prs)
    _title_box(s, "1) Overview")
    _bullets(s, [findings["overview"]] + [f"Internal: {', '.join(findings['stakeholders']['internal'])}", f"External: {', '.join(findings['stakeholders']['external'])}"])

    # 2 Problem
    s = _blank(prs)
    _title_box(s, "2) Problem Statement")
    _bullets(s, _as_lines(findings["problem_statement"], 6))

    # 3 Methodology
    s = _blank(prs)
    _title_box(s, "3) Proposed Methodology")
    _bullets(s, findings["methodology"])

    # 4 Data overview
    s = _blank(prs)
    _title_box(s, "4) Data Overview")
    _bullets(s, _as_lines(findings["data_overview"], 10))

    # 5 Key findings
    s = _blank(prs)
    _title_box(s, "5) Key Findings")
    _bullets(s, findings["key_findings"])

    # 6 Limitations
    s = _blank(prs)
    _title_box(s, "6) Limitations")
    _bullets(s, findings["limitations"])

    # 7 Conclusion
    s = _blank(prs)
    _title_box(s, "7) Conclusion")
    _bullets(s, _as_lines(findings["conclusion"], 6))

    # 8 Recommendations
    s = _blank(prs)
    _title_box(s, "8) Recommendations")
    _bullets(s, findings["recommendations"])

    dest = PRESENTATIONS / f"{findings['code'].replace(' ', '')}_{findings['slug']}.pptx"
    prs.save(dest)
    return dest


def build_html(findings: dict) -> Path:
    figs = []
    for rel in findings.get("figures", []):
        path = ROOT / rel
        if path.exists():
            figs.append(path.relative_to(PRESENTATIONS).as_posix() if False else Path("..", rel).as_posix())

    def section(num, title, body_html):
        return f'<section id="s{num}"><div class="kicker">{num}</div><h2>{html.escape(title)}</h2>{body_html}</section>'

    def lis(items):
        return "<ul>" + "".join(f"<li>{html.escape(str(i))}</li>" for i in items) + "</ul>"

    data_items = [f"{k}: {v}" for k, v in findings["data_overview"].items()]
    img_html = "".join(
        f'<figure><img src="{html.escape(src)}" alt="chart"><figcaption>{html.escape(Path(src).name)}</figcaption></figure>'
        for src in figs
    )
    qa_blocks = []
    for level, rows in findings.get("qa", {}).items():
        qa_blocks.append(f"<h3>{html.escape(level.title())} questions</h3><ol>")
        for row in rows:
            qa_blocks.append(f"<li><strong>{html.escape(row['q'])}</strong><div class='ans'>{html.escape(str(row['a']))}</div></li>")
        qa_blocks.append("</ol>")

    page = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{html.escape(findings['title'])}</title>
  <style>
    :root {{ --navy:#0f172a; --teal:#0e7490; --bg:#f8fafc; --card:#fff; --muted:#475569; }}
    * {{ box-sizing: border-box; }}
    body {{ margin:0; font-family: "Segoe UI", Calibri, sans-serif; background: var(--bg); color: var(--navy); }}
    header {{ background: linear-gradient(135deg, var(--navy), #164e63); color:#fff; padding: 48px 8vw 40px; }}
    header .code {{ letter-spacing:.12em; text-transform:uppercase; color:#67e8f9; font-weight:700; }}
    header h1 {{ margin: 8px 0 12px; font-size: clamp(28px, 4vw, 44px); }}
    nav {{ position: sticky; top:0; background:#fff; border-bottom:1px solid #e2e8f0; display:flex; gap:12px; overflow:auto; padding:10px 8vw; z-index:5; }}
    nav a {{ color: var(--teal); text-decoration:none; font-weight:600; white-space:nowrap; }}
    main {{ padding: 24px 8vw 80px; display:grid; gap:28px; }}
    section {{ background: var(--card); border-radius: 16px; padding: 28px 32px; box-shadow: 0 8px 24px rgba(15,23,42,.06); }}
    .kicker {{ color: var(--teal); font-weight: 800; }}
    h2 {{ margin-top: 4px; }}
    ul {{ line-height: 1.55; }}
    .grid {{ display:grid; grid-template-columns: repeat(auto-fit, minmax(260px,1fr)); gap:16px; }}
    figure {{ margin:0; background:#f1f5f9; border-radius:12px; padding:8px; }}
    img {{ width:100%; border-radius:8px; }}
    figcaption {{ font-size:12px; color:var(--muted); padding:6px; }}
    .ans {{ color:var(--muted); margin:6px 0 14px; font-family: ui-monospace, Consolas, monospace; font-size:13px; white-space:pre-wrap; }}
    footer {{ padding: 24px 8vw; color:#64748b; }}
  </style>
</head>
<body>
  <header>
    <div class="code">{html.escape(findings['code'])}</div>
    <h1>{html.escape(findings['title'])}</h1>
    <p>{html.escape(findings.get('dataset',''))}</p>
  </header>
  <nav>
    <a href="#s1">Overview</a><a href="#s2">Problem</a><a href="#s3">Methodology</a>
    <a href="#s4">Data</a><a href="#s5">Findings</a><a href="#s6">Limitations</a>
    <a href="#s7">Conclusion</a><a href="#s8">Recommendations</a><a href="#qa">Q&amp;A</a>
  </nav>
  <main>
    {section('1', 'Overview', '<p>' + html.escape(findings['overview']) + '</p>' + lis(['Internal: ' + ', '.join(findings['stakeholders']['internal']), 'External: ' + ', '.join(findings['stakeholders']['external'])]))}
    {section('2', 'Problem Statement', '<p>' + html.escape(findings['problem_statement']) + '</p>')}
    {section('3', 'Proposed Methodology', lis(findings['methodology']))}
    {section('4', 'Data Overview', lis(data_items) + '<div class="grid">' + img_html + '</div>')}
    {section('5', 'Key Findings', lis(findings['key_findings']))}
    {section('6', 'Limitations', lis(findings['limitations']))}
    {section('7', 'Conclusion', '<p>' + html.escape(findings['conclusion']) + '</p>')}
    {section('8', 'Recommendations', lis(findings['recommendations']))}
    <section id="qa"><div class="kicker">Appendix</div><h2>Solution guide (assignment questions)</h2>{''.join(qa_blocks)}</section>
  </main>
  <footer>Generated from the cleaned dataset and the assignment brief. Not medical or investment advice.</footer>
</body>
</html>
"""
    dest = PRESENTATIONS / f"{findings['code'].replace(' ', '')}_{findings['slug']}.html"
    dest.write_text(page, encoding="utf-8")
    return dest


def build_all(findings_list: list[dict]) -> list[Path]:
    paths = []
    for findings in findings_list:
        paths.append(build_pptx(findings))
        paths.append(build_html(findings))
    return paths
